import argparse
from pathlib import Path

from docx import Document


def make_synthetic_manuscript(path):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Judgment and Responsibility in Clinical AI", level=0)
    document.add_heading("Abstract", level=1)
    document.add_paragraph(
        "Clinical artificial intelligence can support judgment without replacing "
        "the responsibility to explain a recommendation to a patient."
    )
    document.add_heading("Introduction", level=1)
    document.add_paragraph(
        "Recent debate treats prediction as decision support (Smith, 2024), but "
        "prediction alone cannot settle a normative question."
    )
    document.add_paragraph(
        "This article argues that responsibility remains attached to the clinical "
        "act of recommending and explaining care [1]."
    )
    document.add_heading("Ethical Analysis", level=1)
    document.add_paragraph(
        "The distinction between prediction and judgment protects patient agency "
        "while clarifying professional accountability."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Claim type"
    table.cell(0, 1).text = "Example"
    table.cell(1, 0).text = "Normative"
    table.cell(1, 1).text = "Clinicians remain answerable for recommendations."
    document.add_page_break()
    document.add_heading("Limitations", level=1)
    document.add_paragraph(
        "The argument does not establish the empirical effects of any particular system."
    )
    document.add_heading("Conclusion", level=1)
    document.add_paragraph(
        "AI may inform a choice, but it cannot assume the clinician's relational duty."
    )
    document.add_heading("References", level=1)
    document.add_paragraph(
        "Smith, A. (2024). Clinical judgment and prediction. Journal of Example Ethics. "
        "https://doi.org/10.1234/example.2024.001"
    )
    document.save(target)
    return target


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("output")
    args = parser.parse_args()
    make_synthetic_manuscript(args.output)


if __name__ == "__main__":
    main()
