import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document


DOI_PATTERN = re.compile(r"10\.\d{4,9}/[-._;()/:A-Za-z0-9]+")
AUTHOR_YEAR_PATTERN = re.compile(
    r"\([A-Z][A-Za-z'’-]+(?: et al\.)?,?\s+\d{4}[a-z]?\)"
)
NUMBERED_CITATION_PATTERN = re.compile(r"\[\d+(?:\s*[-,]\s*\d+)*\]")
SENSITIVE_PATTERNS = {
    "email": re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I),
    "phone": re.compile(r"(?<!\d)(?:\+?\d[\d\s-]{7,}\d)(?!\d)"),
    "national_id": re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)"),
    "medical_record_id": re.compile(
        r"\b(?:MRN|medical record(?: number)?)\s*[:#]?\s*[A-Z0-9-]{4,}\b",
        re.I,
    ),
}


class SensitiveDocumentError(ValueError):
    pass


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _document_text(document):
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.extend(cell.text for cell in row.cells)
    return "\n".join(blocks)


def _citations(text):
    found = set()
    for pattern in (DOI_PATTERN, AUTHOR_YEAR_PATTERN, NUMBERED_CITATION_PATTERN):
        found.update(pattern.findall(text))
    return sorted(found)


def _reference_text(document):
    collecting = False
    references = []
    for paragraph in document.paragraphs:
        normalized = paragraph.text.strip().casefold()
        if normalized in {"references", "bibliography"}:
            collecting = True
            continue
        if collecting and paragraph.text.strip():
            references.append(paragraph.text.strip())
    return "\n".join(references)


def _redact(value):
    return value[0] + "***" + value[-1] if len(value) > 2 else "***"


def detect_sensitive_markers(text):
    findings = []
    for label, pattern in SENSITIVE_PATTERNS.items():
        for match in pattern.finditer(text):
            findings.append({"label": label, "match": _redact(match.group(0))})
    return findings


def inspect_docx(path):
    source = Path(path)
    if source.suffix.casefold() != ".docx":
        raise ValueError("input must be a .docx file")
    document = Document(source)
    text = _document_text(document)
    headings = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip() and paragraph.style.name.startswith("Heading")
    ]
    reference_text = _reference_text(document)
    with ZipFile(source) as archive:
        parts = set(archive.namelist())
    return {
        "path": str(source.resolve()),
        "sha256": sha256_file(source),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
        "inline_shape_count": len(document.inline_shapes),
        "headings": headings,
        "citations": _citations(text),
        "reference_fingerprint": hashlib.sha256(
            reference_text.encode("utf-8")
        ).hexdigest(),
        "comments_present": "word/comments.xml" in parts,
        "footnotes_present": "word/footnotes.xml" in parts,
        "sensitive_markers": detect_sensitive_markers(text),
        "text": text,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("--strict", action="store_true")
    args = parser.parse_args()
    manifest = inspect_docx(args.input)
    if args.strict and manifest["sensitive_markers"]:
        labels = sorted({item["label"] for item in manifest["sensitive_markers"]})
        raise SensitiveDocumentError(
            "possible sensitive identifiers detected: " + ", ".join(labels)
        )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
