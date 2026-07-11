import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.apply_revision_spec import RevisionApplicationError, apply_revision_spec
from scripts.build_revision_inventory import build_revision_inventory
from scripts.inspect_docx import inspect_docx, sha256_file
from scripts.make_synthetic_manuscript import make_synthetic_manuscript
from scripts.protect_document import compare_protected_fields


ROOT = Path(__file__).resolve().parents[2]


def revision_for(item, revised_text, confirmation=False):
    return {
        "paragraph_index": item["paragraph_index"],
        "original_text": item["text"],
        "revised_text": revised_text,
        "reason": "Clarify the normative claim.",
        "risk_level": "high" if confirmation else "low",
        "author_confirmation": confirmation,
    }


class ApplyRevisionSpecTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = make_synthetic_manuscript(self.root / "source.docx")
        self.inventory = build_revision_inventory(self.source)
        self.abstract = next(
            item
            for item in self.inventory["paragraphs"]
            if item["text"].startswith("Clinical artificial intelligence")
        )
        self.citation = next(
            item
            for item in self.inventory["paragraphs"]
            if "(Smith, 2024)" in item["text"]
        )

    def tearDown(self):
        self.temp.cleanup()

    def spec(self, revisions):
        return {
            "target_journal": "mhcp",
            "source_sha256": sha256_file(self.source),
            "revisions": revisions,
        }

    def test_applies_exact_safe_revision_without_changing_source(self):
        output = self.root / "revised.docx"
        source_hash = sha256_file(self.source)
        original_style = Document(self.source).paragraphs[
            self.abstract["paragraph_index"]
        ].style.name
        spec = self.spec(
            [
                revision_for(
                    self.abstract,
                    "Clinical AI can support judgment while preserving the duty to explain.",
                )
            ]
        )

        record = apply_revision_spec(self.source, output, spec)

        self.assertEqual(sha256_file(self.source), source_hash)
        self.assertEqual(len(record["applied"]), 1)
        self.assertEqual(record["rejected"], [])
        revised = Document(output)
        self.assertEqual(
            revised.paragraphs[self.abstract["paragraph_index"]].text,
            spec["revisions"][0]["revised_text"],
        )
        self.assertEqual(
            revised.paragraphs[self.abstract["paragraph_index"]].style.name,
            original_style,
        )
        self.assertEqual(
            compare_protected_fields(inspect_docx(self.source), inspect_docx(output)),
            [],
        )

    def test_preflight_rejects_protected_or_mismatched_paragraphs(self):
        output = self.root / "revised.docx"
        mismatch = revision_for(self.abstract, "Revised abstract.")
        mismatch["original_text"] = "Wrong original text."
        protected = revision_for(self.citation, "Citation removed.")

        with self.assertRaises(RevisionApplicationError) as caught:
            apply_revision_spec(self.source, output, self.spec([mismatch, protected]))

        self.assertFalse(output.exists())
        reasons = {item["reason"] for item in caught.exception.record["rejected"]}
        self.assertIn("original_text_mismatch", reasons)
        self.assertIn("protected_paragraph", reasons)

    def test_confirmation_items_are_not_applied(self):
        output = self.root / "revised.docx"
        spec = self.spec(
            [revision_for(self.abstract, "A changed author position.", confirmation=True)]
        )

        record = apply_revision_spec(self.source, output, spec)

        self.assertEqual(record["applied"], [])
        self.assertEqual(len(record["author_confirmation"]), 1)
        self.assertEqual(
            Document(output).paragraphs[self.abstract["paragraph_index"]].text,
            self.abstract["text"],
        )

    def test_existing_output_is_not_overwritten(self):
        output = self.root / "revised.docx"
        output.write_bytes(b"existing")

        with self.assertRaises(FileExistsError):
            apply_revision_spec(
                self.source,
                output,
                self.spec([revision_for(self.abstract, "Revised abstract.")]),
            )

        self.assertEqual(output.read_bytes(), b"existing")

    def test_direct_cli_entrypoint_runs_from_repository_root(self):
        output = self.root / "cli-revised.docx"
        spec_path = self.root / "spec.json"
        spec_path.write_text(
            json.dumps(self.spec([revision_for(self.abstract, "Revised abstract.")])),
            encoding="utf-8",
        )

        result = subprocess.run(
            [
                sys.executable,
                str(ROOT / "scripts" / "apply_revision_spec.py"),
                str(self.source),
                str(output),
                str(spec_path),
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(output.exists())


if __name__ == "__main__":
    unittest.main()
