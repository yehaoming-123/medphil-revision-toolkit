import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from scripts.build_revision_inventory import build_revision_inventory
from scripts.make_synthetic_manuscript import make_synthetic_manuscript


class RevisionInventoryTests(unittest.TestCase):
    def test_synthetic_inventory_protects_citations_and_references(self):
        with tempfile.TemporaryDirectory() as temp:
            source = make_synthetic_manuscript(Path(temp) / "source.docx")

            inventory = build_revision_inventory(source)

            abstract = next(
                item
                for item in inventory["paragraphs"]
                if item["text"].startswith("Clinical artificial intelligence")
            )
            citation = next(
                item
                for item in inventory["paragraphs"]
                if "(Smith, 2024)" in item["text"]
            )
            reference = next(
                item
                for item in inventory["paragraphs"]
                if item["text"].startswith("Smith, A. (2024)")
            )

            self.assertTrue(abstract["eligible"])
            self.assertEqual(abstract["section"], "Abstract")
            self.assertFalse(citation["eligible"])
            self.assertIn("citation", citation["protection_reasons"])
            self.assertFalse(reference["eligible"])
            self.assertIn("reference_section", reference["protection_reasons"])

    def test_complex_paragraphs_are_ineligible(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "complex.docx"
            document = Document()
            document.add_paragraph("Ordinary prose.")

            mixed = document.add_paragraph()
            mixed.add_run("ordinary ")
            mixed.add_run("emphasis").bold = True

            field = document.add_paragraph()
            field_run = field.add_run("field")
            field_run._r.append(OxmlElement("w:fldChar"))

            drawing = document.add_paragraph()
            drawing_run = drawing.add_run("drawing")
            drawing_run._r.append(OxmlElement("w:drawing"))

            hyperlink = document.add_paragraph("link container")
            hyperlink._p.append(OxmlElement("w:hyperlink"))
            document.save(path)

            items = build_revision_inventory(path)["paragraphs"]

            self.assertTrue(items[0]["eligible"])
            self.assertIn("mixed_formatting", items[1]["protection_reasons"])
            self.assertIn("field", items[2]["protection_reasons"])
            self.assertIn("drawing", items[3]["protection_reasons"])
            self.assertIn("hyperlink", items[4]["protection_reasons"])


if __name__ == "__main__":
    unittest.main()
