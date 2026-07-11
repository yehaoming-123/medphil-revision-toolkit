import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.apply_revision_spec import apply_revision_spec
from scripts.build_delivery_bundle import REQUIRED_DELIVERABLES, build_delivery_bundle
from scripts.inspect_docx import inspect_docx, sha256_file
from scripts.protect_document import compare_protected_fields


ROOT = Path(__file__).resolve().parents[2]


def routing_record():
    rows = [
        (1, "mhcp", "医学与哲学的交叉问题清晰"),
        (2, "jmp", "具有临床判断与责任的分析哲学维度"),
        (3, "jme", "包含临床伦理相关性但现实争议不足"),
        (4, "bioethics", "规范问题相关但一般生命伦理范围较弱"),
    ]
    return {
        "ranking": [
            {
                "rank": rank,
                "journal": journal,
                "fit_reasons": [fit],
                "mismatch_reasons": ["基于合成短稿的有限判断"],
                "evidence_limits": ["editorial_inference"],
            }
            for rank, journal, fit in rows
        ],
        "recommended_target": "mhcp",
        "alternative": "jmp",
        "ranking_change_conditions": ["若改为纯分析哲学论证，JMP 可能上升。"],
    }


def fit_record():
    return {
        "target_journal": "mhcp",
        "official_requirements": ["投稿前重新核验当前稿型与声明要求。"],
        "observed_tendencies": ["突出医学实践与哲学问题之间的双向联系。"],
        "abstract_limited": ["本评测未以摘要级外部证据支持具体格式判断。"],
        "editorial_inferences": ["摘要应直接呈现判断与责任的核心区分。"],
        "remaining_risks": ["合成短稿未展开反对意见。"],
        "pre_submission_checks": ["复核引文、参考文献、声明和当前官方要求。"],
    }


class EndToEndRevisionTests(unittest.TestCase):
    def test_synthetic_mhcp_revision_produces_safe_bundle(self):
        source = ROOT / "evals" / "fixtures" / "synthetic_manuscript.docx"
        spec = json.loads(
            (ROOT / "evals" / "fixtures" / "synthetic_revision_spec.json").read_text(
                encoding="utf-8"
            )
        )
        source_hash = sha256_file(source)
        source_document = Document(source)
        frozen_intro = [source_document.paragraphs[index].text for index in (4, 5)]

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            revised = root / "working.docx"
            record = apply_revision_spec(source, revised, spec)
            before = inspect_docx(source)
            after = inspect_docx(revised)

            self.assertEqual(sha256_file(source), source_hash)
            self.assertEqual(len(record["applied"]), 3)
            self.assertEqual(compare_protected_fields(before, after), [])
            revised_document = Document(revised)
            self.assertEqual(
                [revised_document.paragraphs[index].text for index in (4, 5)],
                frozen_intro,
            )
            for revision in spec["revisions"]:
                self.assertEqual(
                    revised_document.paragraphs[revision["paragraph_index"]].text,
                    revision["revised_text"],
                )

            bundle = root / "bundle"
            paths = build_delivery_bundle(
                bundle,
                revised,
                before,
                after,
                record,
                routing_record(),
                fit_record(),
            )
            self.assertEqual(set(paths), REQUIRED_DELIVERABLES)
            self.assertEqual({path.name for path in bundle.iterdir()}, REQUIRED_DELIVERABLES)
            for name in (
                "synthetic_routing_report.md",
                "synthetic_fit_report.md",
                "synthetic_confirmation.md",
            ):
                expected = (ROOT / "evals" / "expected" / name).read_text(
                    encoding="utf-8"
                )
                actual_name = {
                    "synthetic_routing_report.md": "journal_routing_report.md",
                    "synthetic_fit_report.md": "journal_fit_report.md",
                    "synthetic_confirmation.md": "author_confirmation_needed.md",
                }[name]
                self.assertEqual((bundle / actual_name).read_text(encoding="utf-8"), expected)


if __name__ == "__main__":
    unittest.main()
